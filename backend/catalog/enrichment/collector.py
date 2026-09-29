import hashlib
import json
import re
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

from catalog.enrichment.sources import ADAPTERS, CachedClient

EDITION_FIELDS = {
    "isbn",
    "pages",
    "binding",
    "publication_date",
    "publisher",
    "cover_url",
    "cover_source_url",
    "edition",
    "expected_count",
    "completeness",
}


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str))
    temp.replace(path)


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in workbook:
        rows = list(sheet.iter_rows(values_only=True))
        headers = rows[0]
        for row, cells in enumerate(rows[1:], 2):
            if not any(v is not None for v in cells):
                continue
            records.append(
                {
                    "source_file": Path(path).name,
                    "sheet": sheet.title,
                    "row": row,
                    "raw": {str(k): v for k, v in zip(headers, cells, strict=True)},
                }
            )
    workbook.close()
    return records


def claim(value, source, timestamp, status="confirmed"):
    return {
        "value": value,
        "source_url": source["url"],
        "source_label": source.get("label", source["adapter"]),
        "collected_at": timestamp,
        "status": status,
    }


def make_entity(source, parsed, timestamp, raw):
    if source.get("expected_source_id") and parsed.get("source_id") != source["expected_source_id"]:
        raise ValueError("来源版本标识变化：请重新核对版本映射")
    kind = source["kind"]
    fields = {}
    for field, selector in source.get("fields", {}).items():
        if field in EDITION_FIELDS and not source.get("edition_verified"):
            continue
        if isinstance(selector, dict):
            value = selector.get("value")
        else:
            value = parsed.get(selector)
        if value is not None and value != "" and value != []:
            fields[field] = claim(value, source, timestamp)
    if source.get("required_text"):
        text = parsed.get("text", parsed.get("title", ""))
        if not all(token.casefold() in text.casefold() for token in source["required_text"]):
            raise ValueError("来源结构或书目身份变化：核对关键词未匹配")
    entity = {
        "kind": kind,
        "identity_key": source["identity_key"],
        "fields": fields,
        "edition_verified": source.get("edition_verified", False),
        "existing_source_row": source.get("existing_source_row"),
        "members": [],
    }
    entity["archive_source_row"] = source.get("archive_source_row")
    if kind == "series" and source.get("members_verified"):
        excluded = source.get("exclude_members", [])
        included = source.get("include_members")
        members = parsed.get("members", [])
        for position, member in enumerate(members, 1):
            title = member["title"].replace("\\'", "'")
            if title == parsed.get("title") or member.get("source_id") == parsed.get("source_id"):
                continue  # An aggregate audio pack is not a physical child book.
            if source.get("member_pattern") and not re.search(source["member_pattern"], title):
                continue
            if any(word in title for word in excluded) or (included and title not in included):
                continue
            if not member.get("source_id"):
                raise ValueError(f"成员缺少版本来源标识：{title}")
            existing_row = source.get("existing_member_mappings", {}).get(member["source_id"])
            position = source.get("member_positions", {}).get(member["source_id"], position)
            entity["members"].append(
                {
                    "kind": "book",
                    "existing_source_row": existing_row,
                    "identity_key": "xiaobien:book:" + member["source_id"],
                    "edition_verified": True,
                    "position": position,
                    "members": [],
                    "fields": {
                        key: claim(value, source, timestamp)
                        for key, value in {
                            "title": title,
                            "aliases": [raw["书本名称"]]
                            if existing_row == source.get("raw_row")
                            else [],
                            "cover_url": member.get("cover_url", ""),
                            "cover_source_url": source["url"],
                            "language": source.get("member_language", "unknown"),
                            "edition": "小彼恩点读版；分册 ISBN/出版信息待核实",
                            "reading_attributes": {
                                "point_reading": {
                                    "value": "小彼恩；仅适用于此来源的点读版本",
                                    "status": "confirmed",
                                    "source_url": source["url"],
                                }
                            },
                        }.items()
                        if value != "" and value != []
                    },
                }
            )
    return entity


def collect(workbook, mapping, output, cache, refresh=False, progress=None):
    mapping = json.loads(Path(mapping).read_text())
    previous = {}
    if Path(output).exists() and not refresh:
        previous = {
            (r["sheet"], r["row"]): r for r in json.loads(Path(output).read_text())["records"]
        }
    client = CachedClient(cache, refresh=refresh)
    results = []
    for record in read_workbook(workbook):
        entry = mapping["records"].get(str(record["row"]), {})
        fingerprint = hashlib.sha256(
            json.dumps([6, record, entry], sort_keys=True, ensure_ascii=False, default=str).encode()
        ).hexdigest()
        old = previous.get((record["sheet"], record["row"]))
        if old and old.get("fingerprint") == fingerprint:
            results.append(old)
            if progress:
                progress(f"行 {record['row']}：续跑复用")
            continue
        record.update(
            {
                "fingerprint": fingerprint,
                "queries": entry.get("queries", []),
                "candidates": entry.get("candidates", []),
                "issues": entry.get("issues", []),
                "entities": [],
                "status": "pending",
            }
        )
        search_url = (
            "https://www.xiaobien.com/xiaobien-api/diandubi/v1/web/merge/search?"
            + urlencode(
                {
                    "keywords": entry.get(
                        "search_term", str(record["raw"].get("书本名称", "")).strip()
                    ),
                    "per_page": 10,
                    "page": 1,
                }
            )
        )
        discovery = {"url": search_url, "adapter": "xiaobien-search", "confirmed": False}
        for source in [discovery, *entry.get("sources", [])]:
            candidate = {"url": source["url"], "adapter": source["adapter"]}
            try:
                fetched = client.fetch(source["url"])
                parsed = ADAPTERS[source["adapter"]](fetched["text"], source["url"])
                candidate.update(
                    {
                        "title": parsed.get("title", ""),
                        "collected_at": fetched["collected_at"],
                        "parsed": parsed,
                        "status": "collected",
                    }
                )
                # Reviewed mappings specify identity/fields. Scraping alone never confirms a match.
                if source.get("confirmed"):
                    record["entities"].append(
                        make_entity(source, parsed, fetched["collected_at"], record["raw"])
                    )
            except (
                ValueError,
                KeyError,
                TypeError,
                OSError,
                IndexError,
                AttributeError,
                RecursionError,
            ) as exc:
                candidate.update({"status": "failed", "reason": str(exc)})
                record["issues"].append(f"采集失败 {source['url']}: {exc}")
            record["candidates"].append(candidate)
        record["status"] = "ready" if record["entities"] else "pending"
        if not record["candidates"]:
            record["issues"].append("当前公开检索未取得可核实版本；需书影/版权页或具体套装目录")
        results.append(record)
        write_json(output, {"schema_version": 1, "records": results})
        if progress:
            progress(
                f"行 {record['row']}：{record['status']}，{len(record['entities'])} 个可靠对象"
            )
    write_json(output, {"schema_version": 1, "records": results})
    return results
